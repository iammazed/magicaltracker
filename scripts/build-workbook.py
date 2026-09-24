"""
Builds data/catalog-entry.xlsx FROM data/*.csv.

The workbook is a DERIVED artifact, never a source. It is regenerated from the
CSVs on demand, which is what stops the two drifting apart — previously the
generator produced an empty template, so regenerating it silently discarded
every row.

Run it through `npm run data:workbook`, which exports the vocabularies from
scripts/vocabulary.mjs first so the dropdowns cannot disagree with the
validator.

Requires openpyxl:  pip install openpyxl
"""

import csv
import io
import json
import os
import sys

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
VOCAB = json.load(open(sys.argv[1], encoding="utf-8"))
OUT = os.path.join(ROOT, "data", "catalog-entry.xlsx")

TEAL, NAVY, VIOLET = "00807E", "2B4063", "512663"
GOLD_F, RED_F, RULE = "FBF3E0", "FBECE9", "D9E3E3"

FONT = "Arial"
hdr_font = Font(name=FONT, bold=True, color="FFFFFF", size=10)
body_font = Font(name=FONT, size=10)
thin = Side(style="thin", color=RULE)
border = Border(left=thin, right=thin, top=thin, bottom=thin)

MAX_ROW = 600


def read_csv(name, expected):
    """Returns (rows, note). Missing or mis-headered files yield no rows."""
    path = os.path.join(ROOT, "data", name)
    if not os.path.exists(path):
        return [], "not found"
    with io.open(path, encoding="utf-8-sig", newline="") as f:
        rows = list(csv.reader(f))
    if not rows:
        return [], "empty"
    header = [h.strip() for h in rows[0]]
    if header != expected:
        missing = [c for c in expected if c not in header]
        extra = [c for c in header if c not in expected]
        bits = []
        if missing:
            bits.append("missing " + ", ".join(missing))
        if extra:
            bits.append("unexpected " + ", ".join(extra))
        return [], "HEADER MISMATCH: " + "; ".join(bits)
    # Re-key each row to the expected order so a reordered CSV still lands right.
    idx = {h: i for i, h in enumerate(header)}
    out = [[(r[idx[c]] if idx[c] < len(r) else "") for c in expected] for r in rows[1:]]
    return out, f"{len(out)} rows"


wb = Workbook()

# ── Lists ────────────────────────────────────────────────────────────────
ls = wb.active
ls.title = "Lists"
ls.sheet_properties.tabColor = NAVY

areas = [a["id"] for a in VOCAB["areas"]]
COLS = {
    "areas": areas,
    "venue_kind": VOCAB["kinds"],
    "service_type": VOCAB["service"],
    "dining_style": VOCAB["styles"],
    "status": VOCAB["status"],
    "price_tier": ["1", "2", "3", "4"],
    "yesno": ["TRUE", "FALSE"],
    "resort_tier": VOCAB["tiers"],
    "destination": ["wdw"],
    "transport": VOCAB["transport"],
    "tags": VOCAB["tags"],
}

col_i = 1
named = {}
for name, values in COLS.items():
    L = get_column_letter(col_i)
    ls.cell(row=1, column=col_i, value=name).font = Font(name=FONT, bold=True, size=9, color=TEAL)
    for r, v in enumerate(values, start=2):
        ls.cell(row=r, column=col_i, value=v).font = body_font
    named[name] = f"Lists!${L}$2:${L}${1 + len(values)}"
    ls.column_dimensions[L].width = max(14, len(name) + 4)
    col_i += 1

# One named range per area, for the dependent sub_area dropdown. The name is
# the area id with hyphens stripped, matching SUBSTITUTE(...,"-","") at the
# point of use.
for area in areas:
    L = get_column_letter(col_i)
    vals = VOCAB["subAreas"].get(area, [])
    ls.cell(row=1, column=col_i, value=area).font = Font(name=FONT, bold=True, size=9, color=VIOLET)
    for r, v in enumerate(vals, start=2):
        ls.cell(row=r, column=col_i, value=v).font = body_font
    # Areas with no sub-areas get one blank cell so INDIRECT resolves instead
    # of erroring; the dropdown simply offers one empty choice.
    named[area.replace("-", "")] = f"Lists!${L}$2:${L}${1 + max(len(vals), 1)}"
    ls.column_dimensions[L].width = 20
    col_i += 1

for nm, ref in named.items():
    wb.defined_names.add(DefinedName(nm, attr_text=ref))
ls.freeze_panes = "A2"


# ── Sheet builder ────────────────────────────────────────────────────────
def build(title, headers, widths, colour, rows, dropdowns, comments):
    ws = wb.create_sheet(title)
    ws.sheet_properties.tabColor = colour
    fill = PatternFill("solid", fgColor=colour)
    for i, h in enumerate(headers, start=1):
        c = ws.cell(row=1, column=i, value=h)
        c.font, c.fill, c.border = hdr_font, fill, border
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = widths.get(h, 16)
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "A2"

    ix = {h: i + 1 for i, h in enumerate(headers)}
    letter = lambda h: get_column_letter(ix[h])

    for col, formula, stop, prompt in dropdowns:
        dv = DataValidation(
            type="list", formula1=formula, allow_blank=True,
            showErrorMessage=True, errorStyle="stop" if stop else "warning",
            error="Not a valid value. Pick from the dropdown.",
            errorTitle="Invalid value",
            showInputMessage=bool(prompt), prompt=prompt, promptTitle="",
        )
        # openpyxl's showDropDown maps to the OOXML "suppress" flag, so False
        # is what makes the arrow appear.
        dv.showDropDown = False
        ws.add_data_validation(dv)
        dv.add(f"{letter(col)}2:{letter(col)}{MAX_ROW}")

    for col, text in comments.items():
        ws[f"{letter(col)}1"].comment = Comment(text, "MagicalTracker", height=170, width=290)

    for r, row in enumerate(rows, start=2):
        for c, val in enumerate(row, start=1):
            ws.cell(row=r, column=c, value=val).font = body_font

    # Rows with an id but no verified_on tint red, so progress is visible.
    ws.conditional_formatting.add(
        f"A2:{get_column_letter(len(headers))}{MAX_ROW}",
        FormulaRule(formula=[f'AND($A2<>"",${letter("verified_on")}2="")'],
                    fill=PatternFill("solid", fgColor=RED_F), stopIfTrue=False))
    return ws


vh, rh = VOCAB["venueCols"], VOCAB["resortCols"]
venue_rows, venue_note = read_csv("venues.csv", vh)
resort_rows, resort_note = read_csv("resorts.csv", rh)

GATE = ("THE GATE. Format YYYY-MM-DD.\n\nFill this in ONLY after reading the official "
        "disneyworld.com page. Rows without it are refused by npm run db:import.")

build("Venues", vh, {
    "id": 24, "name": 32, "destination_id": 13, "area_id": 19, "sub_area": 19,
    "resort_id": 22, "venue_kind": 13, "service_type": 13, "dining_style": 13,
    "cuisine": 16, "price_tier": 9, "accepts_reservations": 13,
    "is_character_dining": 13, "is_signature": 11, "status": 13,
    "lat": 12, "lng": 12, "menu_url": 34, "description": 46, "tags": 24,
    "verified_on": 12, "source_url": 34,
}, TEAL, venue_rows, [
    ("destination_id", "=destination", True, None),
    ("area_id", "=areas", True, None),
    ("sub_area", '=INDIRECT(SUBSTITUTE($E2,"-",""))', False,
     "Choices depend on the area picked in this row. Leave blank for resort venues."),
    ("venue_kind", "=venue_kind", True, None),
    ("service_type", "=service_type", True, None),
    ("dining_style", "=dining_style", True, None),
    ("status", "=status", True, None),
    ("price_tier", "=price_tier", True, None),
    ("accepts_reservations", "=yesno", True, None),
    ("is_character_dining", "=yesno", True, None),
    ("is_signature", "=yesno", True, None),
], {
    "tags": "Pipe-separated, no spaces:  lunch|dinner\n\nAllowed:\n  " + "\n  ".join(VOCAB["tags"]),
    "lat": ("Drop a pin in Google Maps at the actual building, right-click, copy "
            "coordinates.\n\nMust fall inside 28.28 to 28.44. A geocoding API returns "
            "the park entrance, not the venue."),
    "description": ("One or two sentences IN YOUR OWN WORDS.\n\nNever paste Disney "
                    "marketing copy - that is a copyright issue separate from trademark."),
    "verified_on": GATE,
})

build("Resorts", rh, {
    "id": 26, "name": 38, "destination_id": 13, "area_id": 22, "tier": 12,
    "transport": 24, "transport_notes": 42, "lat": 12, "lng": 12,
    "official_url": 34, "description": 46, "status": 13, "verified_on": 12,
    "source_url": 34,
}, VIOLET, resort_rows, [
    ("destination_id", "=destination", True, None),
    ("area_id", "=areas", True,
     "Resort areas only: mk-resort-area, epcot-resort-area, ak-resort-area, "
     "springs-resort-area, sports-resort-area"),
    ("tier", "=resort_tier", True, None),
    ("status", "=status", True, None),
], {
    "transport": ("Pipe-separated, no spaces:  monorail|boat|bus\n\nAllowed:\n  "
                  + "\n  ".join(VOCAB["transport"])
                  + "\n\nDrives the Transportation Challenge - a wrong value makes a "
                    "badge unearnable."),
    "transport_notes": ('Optional free text: what the modes actually connect to, e.g. '
                        '"Walk or boat to EPCOT and Hollywood Studios".\n\nFill only '
                        'where it is not obvious. A value resort whose only mode is bus '
                        'needs nothing here.'),
    "verified_on": GATE,
})

# ── Read me ──────────────────────────────────────────────────────────────
ref = wb.create_sheet("Read me first", 0)
ref.sheet_properties.tabColor = TEAL
ref.column_dimensions["A"].width = 4
ref.column_dimensions["B"].width = 26
ref.column_dimensions["C"].width = 98
ref.sheet_view.showGridLines = False

LINES = [
    ("MagicalTracker", "Catalog data entry workbook", "h1"),
    ("", "", ""),
    ("THIS FILE IS GENERATED", "Built from data/venues.csv and data/resorts.csv by "
     "`npm run data:workbook`. It is a copy, not the original. Regenerating it "
     "overwrites anything typed here that has not been exported back to CSV.", "h2"),
    ("", "", ""),
    ("The round trip", "", "h2"),
    ("1.", "npm run data:workbook            rebuild this file from the current CSVs", ""),
    ("2.", "Upload to Google Sheets          File > Import > Replace spreadsheet", ""),
    ("3.", "Edit there", ""),
    ("4.", "File > Download > CSV            once per tab, into data/", ""),
    ("5.", "npm run data:normalize           repair the dates Sheets reformats", ""),
    ("6.", "npm run validate:data            catch typos and broken references", ""),
    ("7.", "npm run db:import                push verified rows to Supabase", ""),
    ("", "", ""),
    ("The one rule", "", "h2"),
    ("verified_on", "Fill this in ONLY after reading the official disneyworld.com page. "
     "Rows without it are skipped by the importer and can never reach the app. Rows "
     "missing it tint red.", ""),
    ("Never guess", "Do not fill rows from memory, from a fan wiki, or from an AI. Closed "
     "restaurants linger on fan sites for years, and invented coordinates look completely "
     "plausible.", ""),
    ("Your own words", "Write descriptions yourself. Pasting Disney marketing copy is a "
     "copyright issue separate from trademark.", ""),
    ("", "", ""),
    ("Dropdowns", "", "h2"),
    ("Most columns", "Click the cell and use the arrow. Invalid values are rejected.", ""),
    ("sub_area", "Its choices change based on the area_id in that row, so pick the area "
     "first. Uses INDIRECT, which Google Sheets may drop on import - if it does, the "
     "validator still catches any mismatch.", ""),
    ("tags / transport", "Free text, pipe-separated, no spaces:  lunch|dinner", ""),
    ("", "", ""),
    ("Current contents", "", "h2"),
    ("Venues", venue_note, ""),
    ("Resorts", resort_note, ""),
]

for i, (b, c, style) in enumerate(LINES, start=2):
    cb, cc = ref.cell(row=i, column=2, value=b), ref.cell(row=i, column=3, value=c)
    if style == "h1":
        cb.font = Font(name=FONT, bold=True, size=14, color=TEAL)
        cc.font = Font(name=FONT, size=11, color=NAVY)
    elif style == "h2":
        cb.font = Font(name=FONT, bold=True, size=11, color=NAVY)
        cc.font = body_font
    else:
        cb.font = Font(name=FONT, bold=True, size=10)
        cc.font = body_font
    cc.alignment = Alignment(wrap_text=True, vertical="top")

wb.save(OUT)
print(f"  venues   {venue_note}")
print(f"  resorts  {resort_note}")
print(f"  wrote    data/catalog-entry.xlsx")
if "MISMATCH" in venue_note or "MISMATCH" in resort_note:
    sys.exit(1)
